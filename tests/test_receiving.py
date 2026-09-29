"""Reference-lab receiving: Prevencio boxes, sample intake, QA / clinical review / rejection, spreadsheet export."""
import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from conftest import uniq, xlsx_bytes


@pytest.fixture
def prevencio(app, master):
    with app.app.app_context():
        c = app.ReferenceClient.query.filter_by(code='PREV').one()
        return c.id


@pytest.fixture
def lab(login_as):
    return {'staff': login_as('staff'), 'director': login_as('director')}


def import_roster(app, director, client_id):
    # Headerless "Customers" tab like the real spreadsheet: name, account ID# (leading zeros, one shared ID).
    data = xlsx_bytes([['Harbor Heart Clinic', '071'], ['Bay Cardiology', '114'], ['Bay Cardiology North', '114'], ['Valley Internal Medicine', 2057.0]], sheet='Customers')
    r = director.post(f'/receiving/{client_id}', {'action': 'import_accounts', 'roster': (io.BytesIO(data), 'roster.xlsx')}, follow_redirects=True)
    assert b'Account list imported' in r.data
    with app.app.app_context():
        # standing order: Harbor gets CADhs by default
        c = app.Clinic.query.filter_by(client_id=client_id, account_id='071').one()
        if not app.StandingOrder.query.filter_by(clinic_id=c.id).first():
            app.db.session.add(app.StandingOrder(clinic_id=c.id, hart_cadhs=True, hart_cve=False, active=True))
            app.db.session.commit()


def open_box(app, staff, client_id, tracking='9622001900005105596800123456789012'):
    r = staff.post(f'/receiving/{client_id}', {'action': 'receive_box', 'tracking_no': tracking})
    assert r.status_code == 302 and '/receiving/boxes/' in r.location
    return int(r.location.rstrip('/').split('/')[-1])


def sample_form(**kw):
    today = kw.pop('today', None)
    data = {'action': 'add_sample', 'order_code': uniq('SAM'), 'collected': today or '', 'account_id': '071', 'account_name': '',
            'last_name': 'Doe', 'first_name': 'Jane', 'dob': '1961-04-02', 'sex': 'F', 'test_HART-CADHS': '1', 'decision': 'accept'}
    data.update(kw)
    return data


def today_iso(app):
    return app.utcnow().date().isoformat()


def test_prevencio_is_seeded(app, prevencio):
    with app.app.app_context():
        c = app.db.session.get(app.ReferenceClient, prevencio)
        assert c.stability_hours == 72
        assert [x['code'] for x in c.test_columns] == ['HART-CADHS', 'HART-CVE']


def test_roster_import_handles_leading_zeros_and_shared_ids(app, lab, prevencio):
    import_roster(app, lab['director'], prevencio)
    with app.app.app_context():
        ids = sorted(c.account_id for c in app.Clinic.query.filter_by(client_id=prevencio).all())
        assert '071' in ids and '2057' in ids and ids.count('114') == 2
    # re-import is idempotent
    data = xlsx_bytes([['Harbor Heart Clinic', '071']], sheet='Customers')
    r = lab['director'].post(f'/receiving/{prevencio}', {'action': 'import_accounts', 'roster': (io.BytesIO(data), 'r.xlsx')}, follow_redirects=True)
    assert b'0 new, 1 already known' in r.data


def test_box_and_sample_intake_happy_path(app, lab, prevencio):
    staff = lab['staff']
    import_roster(app, lab['director'], prevencio)
    box = open_box(app, staff, prevencio)
    with app.app.app_context():
        b = app.db.session.get(app.Shipment, box)
        assert b.tracking_no == '123456789012' and b.received_by == staff.user['id'] and b.status == 'Open'
    # scanning the same label again reopens the same open box
    assert open_box(app, staff, prevencio) == box

    form = sample_form(collected=today_iso(app), **{'test_HART-CVE': '1'})
    r = staff.post(f'/receiving/boxes/{box}', form, follow_redirects=True)
    assert b'added.' in r.data
    with app.app.app_context():
        o = app.Order.query.filter_by(accession_no=form['order_code']).one()
        assert o.status == 'Received' and o.shipment_id == box and o.patient_last_name == 'Doe' and o.patient_sex == 'Female'
        assert app.db.session.get(app.Clinic, o.clinic_id).account_id == '071'
        codes = {app.db.session.get(app.Test, ot.test_id).code for ot in app.OrderTest.query.filter_by(order_id=o.id)}
        assert codes == {'HART-CADHS', 'HART-CVE'}
        assert app.Sample.query.filter_by(sample_no=form['order_code']).one().collected_at == today_iso(app)
        assert app.SampleException.query.filter_by(order_id=o.id).count() == 0
        oid = o.id
    # accepted samples flow into the normal lab workflow
    staff.post(f'/orders/{oid}', {'action': 'testing'})
    with app.app.app_context():
        assert app.db.session.get(app.Order, oid).status == 'Testing'
    assert staff.get(f'/receiving/boxes/{box}').status_code == 200


def test_intake_validation(app, lab, prevencio):
    staff = lab['staff']
    import_roster(app, lab['director'], prevencio)
    box = open_box(app, staff, prevencio, tracking=uniq('7'))
    first = sample_form(collected=today_iso(app))
    staff.post(f'/receiving/boxes/{box}', first)
    tomorrow = (app.utcnow().date() + timedelta(days=1)).isoformat()
    bad = [
        sample_form(order_code=first['order_code']),               # duplicate Order ID
        sample_form(last_name=''),                                  # missing name
        sample_form(dob=''),                                        # missing DOB
        sample_form(dob=tomorrow),                                  # DOB in future
        sample_form(collected=tomorrow),                            # collected after receipt
        sample_form(dob='31/31/1990'),                              # unparseable date
        sample_form(account_id='999'),                              # unknown account, no name
        {**sample_form(), 'test_HART-CADHS': ''},                   # no tests
        sample_form(decision='qa', reason=''),                      # hold without reason
    ]
    with app.app.app_context():
        before = app.Order.query.count()
    for data in bad:
        r = staff.post(f'/receiving/boxes/{box}', data)
        assert r.status_code == 200, data
    with app.app.app_context():
        assert app.Order.query.count() == before
    # unknown account WITH a name is added to Prevencio's list
    new = sample_form(account_id='555', account_name='Brand New Practice', collected=today_iso(app))
    staff.post(f'/receiving/boxes/{box}', new)
    with app.app.app_context():
        c = app.Clinic.query.filter_by(client_id=prevencio, account_id='555').one()
        assert c.name == 'Brand New Practice'


def test_shared_account_id_needs_name(app, lab, prevencio):
    staff = lab['staff']
    import_roster(app, lab['director'], prevencio)
    box = open_box(app, staff, prevencio, tracking=uniq('8'))
    f = sample_form(account_id='114', account_name='Bay Cardiology North', collected=today_iso(app))
    staff.post(f'/receiving/boxes/{box}', f)
    with app.app.app_context():
        o = app.Order.query.filter_by(accession_no=f['order_code']).one()
        assert app.db.session.get(app.Clinic, o.clinic_id).name == 'Bay Cardiology North'


def test_stability_and_missing_date_go_to_qa_and_block_testing(app, lab, prevencio):
    staff = lab['staff']
    import_roster(app, lab['director'], prevencio)
    box = open_box(app, staff, prevencio, tracking=uniq('6'))
    old = (app.utcnow().date() - timedelta(days=5)).isoformat()
    stale, undated = sample_form(collected=old), sample_form(collected='')
    staff.post(f'/receiving/boxes/{box}', stale)
    staff.post(f'/receiving/boxes/{box}', undated)
    with app.app.app_context():
        o1 = app.Order.query.filter_by(accession_no=stale['order_code']).one()
        o2 = app.Order.query.filter_by(accession_no=undated['order_code']).one()
        e1 = app.SampleException.query.filter_by(order_id=o1.id).one()
        e2 = app.SampleException.query.filter_by(order_id=o2.id).one()
        assert (e1.kind, e1.reason, e1.status) == ('QA', 'Stability >72 hours', 'Open')
        assert (e2.kind, e2.reason) == ('QA', 'No sample collection date')
        oid1, ex1, oid2, ex2 = o1.id, e1.id, o2.id, e2.id
    # on hold: testing is blocked
    staff.post(f'/orders/{oid1}', {'action': 'testing'})
    with app.app.app_context():
        assert app.db.session.get(app.Order, oid1).status == 'Received'
    qa_page = staff.get(f'/receiving/{prevencio}/exceptions?kind=QA')
    assert stale['order_code'].encode() in qa_page.data
    # resolve needs a note; then testing may proceed
    staff.post(f'/receiving/{prevencio}/exceptions?kind=QA', {'exception_id': str(ex1), 'action': 'resolve', 'resolution': ''})
    with app.app.app_context():
        assert app.db.session.get(app.SampleException, ex1).status == 'Open'
    staff.post(f'/receiving/{prevencio}/exceptions?kind=QA', {'exception_id': str(ex1), 'action': 'resolve', 'resolution': 'Director approved, frozen on arrival'})
    staff.post(f'/orders/{oid1}', {'action': 'testing'})
    with app.app.app_context():
        assert app.db.session.get(app.Order, oid1).status == 'Testing'
    # the other goes to clinical review, then is rejected
    staff.post(f'/receiving/{prevencio}/exceptions?kind=QA', {'exception_id': str(ex2), 'action': 'to_clinical', 'resolution': 'Ask provider for date'})
    with app.app.app_context():
        clin = app.SampleException.query.filter_by(order_id=oid2, kind='CLINICAL').one()
        cid = clin.id
    staff.post(f'/receiving/{prevencio}/exceptions?kind=CLINICAL', {'exception_id': str(cid), 'action': 'reject', 'resolution': 'Provider could not confirm date'})
    with app.app.app_context():
        assert app.db.session.get(app.Order, oid2).status == 'Rejected'
        assert app.SampleException.query.filter_by(order_id=oid2, kind='REJECTED').count() == 1
    # rejected orders cannot be processed
    staff.post(f'/orders/{oid2}', {'action': 'testing'})
    with app.app.app_context():
        assert app.db.session.get(app.Order, oid2).status == 'Rejected'
    assert staff.get(f'/receiving/{prevencio}/exceptions?kind=REJECTED').status_code == 200


def test_reject_at_intake_and_remove_and_close(app, lab, prevencio):
    staff = lab['staff']
    import_roster(app, lab['director'], prevencio)
    box = open_box(app, staff, prevencio, tracking=uniq('5'))
    rej = {**sample_form(collected=today_iso(app), decision='reject', reason='Leaked / broken tube'), 'test_HART-CADHS': ''}
    staff.post(f'/receiving/boxes/{box}', rej)
    mistake = sample_form(collected=today_iso(app))
    staff.post(f'/receiving/boxes/{box}', mistake)
    with app.app.app_context():
        assert app.Order.query.filter_by(accession_no=rej['order_code']).one().status == 'Rejected'
        mid = app.Order.query.filter_by(accession_no=mistake['order_code']).one().id
    staff.post(f'/receiving/boxes/{box}', {'action': 'remove_sample', 'order_id': str(mid)})
    with app.app.app_context():
        assert app.db.session.get(app.Order, mid) is None
        assert app.Sample.query.filter_by(sample_no=mistake['order_code']).count() == 0
    staff.post(f'/receiving/boxes/{box}', {'action': 'close'})
    late = sample_form(collected=today_iso(app))
    staff.post(f'/receiving/boxes/{box}', late)
    with app.app.app_context():
        assert app.db.session.get(app.Shipment, box).status == 'Closed'
        assert app.Order.query.filter_by(accession_no=late['order_code']).count() == 0
    staff.post(f'/receiving/boxes/{box}', {'action': 'reopen'})
    with app.app.app_context():
        assert app.db.session.get(app.Shipment, box).status == 'Open'


def test_export_matches_spreadsheet_layout(app, lab, prevencio):
    staff = lab['staff']
    import_roster(app, lab['director'], prevencio)
    box = open_box(app, staff, prevencio, tracking=uniq('4'))
    good = sample_form(collected=today_iso(app), **{'test_HART-CVE': '1'})
    stale = sample_form(collected=(app.utcnow().date() - timedelta(days=4)).isoformat())
    staff.post(f'/receiving/boxes/{box}', good)
    staff.post(f'/receiving/boxes/{box}', stale)
    month = app.utcnow().strftime('%Y-%m')
    r = staff.get(f'/receiving/{prevencio}/export.xlsx?month={month}')
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames[1:] == ['Quality Assurance (QA)', 'Lab Clinical Review', 'Rejected', 'Customers']
    main = wb.worksheets[0]
    header = [c.value for c in main[1]]
    assert header == ['Order ID', 'Date Received', 'Date Collected', 'Account ID#', 'Account Name', 'Patient Last Name', 'Patient First Name',
                      'Patient DOB', 'Sex', 'CADhs', 'CVE', '# of Tests', 'Received By', 'FedEx #', 'Notes']
    rows = {row[0]: row for row in main.iter_rows(min_row=2, values_only=True)}
    g = rows[good['order_code']]
    assert g[3] == '071' and g[5] == 'Doe' and g[8] == 'F' and g[9:12] == (1, 1, 2)
    assert g[12] == 'Test' and g[13]  # received-by first name, FedEx #
    qa_ids = [row[0] for row in wb['Quality Assurance (QA)'].iter_rows(min_row=2, values_only=True)]
    assert stale['order_code'] in qa_ids and good['order_code'] not in qa_ids
    assert staff.get(f'/receiving/{prevencio}/export.xlsx?month=bad').status_code == 400


def test_client_lab_admin_and_access(app, login_as, prevencio):
    staff, director, doctor = login_as('staff'), login_as('director'), login_as('customer')
    for path in ('/receiving', f'/receiving/{prevencio}', f'/receiving/{prevencio}/exceptions', f'/receiving/{prevencio}/export.xlsx'):
        assert doctor.get(path).status_code == 403, path
    assert staff.post('/receiving', {'name': 'Other Lab', 'code': 'OTH'}).status_code == 403
    assert staff.post(f'/receiving/{prevencio}', {'action': 'settings', 'stability_hours': '1'}).status_code == 403
    r = director.post('/receiving', {'name': uniq('Partner Lab '), 'code': uniq('PL')})
    assert r.status_code == 302
    director.post(f'/receiving/{prevencio}', {'action': 'settings', 'stability_hours': '96', 'active': '1',
                                            'col_label': ['CADhs', 'CVE'], 'col_code': ['HART-CADHS', 'HART-CVE']})
    with app.app.app_context():
        assert app.db.session.get(app.ReferenceClient, prevencio).stability_hours == 96
    director.post(f'/receiving/{prevencio}', {'action': 'settings', 'stability_hours': '72', 'active': '1',
                                            'col_label': ['CADhs', 'CVE'], 'col_code': ['HART-CADHS', 'HART-CVE']})
    assert director.post(f'/receiving/{prevencio}', {'action': 'settings', 'col_label': ['X'], 'col_code': ['NOT-A-TEST']}).status_code == 400
    assert staff.get('/receiving').status_code == 200


def test_tracking_normalisation(app, master):
    with app.app.app_context():
        assert app.normalize_tracking('9622001900005105596800 773456789012') == '773456789012'
        assert app.normalize_tracking('7734 5678 9012,') == '773456789012'
        assert app.normalize_tracking('') == ''


def test_shared_account_id_without_name_is_refused(app, lab, prevencio):
    staff = lab['staff']
    import_roster(app, lab['director'], prevencio)
    box = open_box(app, staff, prevencio, tracking=uniq('3'))
    f = sample_form(account_id='114', account_name='', collected=today_iso(app))
    r = staff.post(f'/receiving/boxes/{box}', f)
    assert b'shared by several accounts' in r.data
    with app.app.app_context():
        assert app.Order.query.filter_by(accession_no=f['order_code']).count() == 0


def test_delete_box_only_when_empty(app, lab, prevencio):
    staff = lab['staff']
    import_roster(app, lab['director'], prevencio)
    # create both first: SQLite reuses the id of a just-deleted row
    empty = open_box(app, staff, prevencio, tracking=uniq('2'))
    full = open_box(app, staff, prevencio, tracking=uniq('1'))
    staff.post(f'/receiving/boxes/{full}', sample_form(collected=today_iso(app)))
    staff.post(f'/receiving/boxes/{empty}', {'action': 'delete_box'})
    r = staff.post(f'/receiving/boxes/{full}', {'action': 'delete_box'}, follow_redirects=True)
    assert b'Only an empty box can be deleted' in r.data
    with app.app.app_context():
        assert app.db.session.get(app.Shipment, empty) is None
        assert app.db.session.get(app.Shipment, full) is not None
        assert app.Audit.query.filter_by(action='BOX_DELETED', entity_id=empty).count() == 1


def _pdf(pages=2):
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for i in range(pages):
        c.drawString(72, 760, f'Test requisition page {i + 1}')
        c.showPage()
    c.save()
    return buf.getvalue()


def _png():
    from PIL import Image
    out = io.BytesIO()
    Image.new('RGB', (300, 400), 'white').save(out, 'PNG')
    return out.getvalue()


def test_fill_automatically_upload_and_review(app, lab, prevencio, login_as):
    staff = lab['staff']
    import_roster(app, lab['director'], prevencio)
    box = open_box(app, staff, prevencio, tracking=uniq('9'))
    assert b'Fill automatically' in staff.get(f'/receiving/boxes/{box}').data
    r = staff.post(f'/receiving/boxes/{box}', {'action': 'upload_docs',
                   'documents': [(io.BytesIO(_pdf(2)), 'box-scan.pdf'), (io.BytesIO(_png()), 'photo.png'), (io.BytesIO(b'not a pdf'), 'notes.txt')]},
                   follow_redirects=True)
    assert b'Uploaded 3 pages' in r.data and b'upload a PDF or an image' in r.data
    with app.app.app_context():
        drafts = app.IntakeDraft.query.filter_by(shipment_id=box).order_by(app.IntakeDraft.id).all()
        assert [d.page_no for d in drafts] == [0, 1, 0] and all(d.status == 'Pending' for d in drafts)
        d1, d2, d3 = (d.id for d in drafts)
    # page image renders for lab staff only
    img = staff.get(f'/receiving/drafts/{d1}/page.png')
    assert img.status_code == 200 and img.data[:8] == b'\x89PNG\r\n\x1a\n'
    assert login_as('customer').get(f'/receiving/drafts/{d1}/page.png').status_code == 403
    # reviewing a page shows it next to the form and carries the draft id
    page = staff.get(f'/receiving/boxes/{box}?mode=auto&draft={d1}').get_data(as_text=True)
    assert f'name="draft_id" value="{d1}"' in page and f'/receiving/drafts/{d1}/page.png' in page
    # adding the sample marks the page done and moves on to the next page
    f = {**sample_form(collected=today_iso(app)), 'draft_id': str(d1)}
    r = staff.post(f'/receiving/boxes/{box}', f)
    assert r.status_code == 302 and f'draft={d2}' in r.location
    with app.app.app_context():
        dd = app.db.session.get(app.IntakeDraft, d1)
        o = app.Order.query.filter_by(accession_no=f['order_code']).one()
        assert dd.status == 'Added' and dd.order_id == o.id
        oid = o.id
    # skip a non-requisition page
    staff.post(f'/receiving/boxes/{box}', {'action': 'discard_draft', 'draft_id': str(d3)})
    with app.app.app_context():
        assert app.db.session.get(app.IntakeDraft, d3).status == 'Discarded'
    # removing the sample puts its page back in the review list
    staff.post(f'/receiving/boxes/{box}', {'action': 'remove_sample', 'order_id': str(oid)})
    with app.app.app_context():
        assert app.db.session.get(app.IntakeDraft, d1).status == 'Pending'
    # closed boxes do not accept uploads
    staff.post(f'/receiving/boxes/{box}', {'action': 'close'})
    r = staff.post(f'/receiving/boxes/{box}', {'action': 'upload_docs', 'documents': [(io.BytesIO(_pdf(1)), 'late.pdf')]}, follow_redirects=True)
    assert b'Reopen the box' in r.data
    # deleting the (now empty) box removes its paperwork too
    staff.post(f'/receiving/boxes/{box}', {'action': 'reopen'})
    staff.post(f'/receiving/boxes/{box}', {'action': 'delete_box'})
    with app.app.app_context():
        assert app.db.session.get(app.Shipment, box) is None
        assert app.IntakeDraft.query.filter_by(shipment_id=box).count() == 0
        assert app.ShipmentDocument.query.filter_by(shipment_id=box).count() == 0


def test_extract_fields_placeholder_returns_nothing(app, prevencio):
    with app.app.app_context():
        assert app.extract_fields(None, app.db.session.get(app.ReferenceClient, prevencio)) == {}
